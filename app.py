"""NotebookLM PDF -> editable PowerPoint.

NotebookLM PDFs are flat pictures (no text layer). This app:
  1. renders each page to an image,
  2. reads the text (embedded text first, OCR otherwise),
  3. erases that text from the image (so no double text),
  4. places native, editable PowerPoint text boxes on top.
Charts/icons/illustrations stay part of the background picture.
"""
import io
import os
import shutil
from pathlib import Path

import fitz  # PyMuPDF
import numpy as np
import streamlit as st
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

try:
    import cv2  # optional, gives cleaner background repair on gradients/photos
except ImportError:
    cv2 = None

APP_TITLE = "NotebookLM PDF → Editable PowerPoint"
OCR_FACTOR = 0.8  # OCR line height -> approximate font size


# ----------------------------------------------------------------- helpers
def find_tessdata():
    cands = [os.environ.get("TESSDATA_PREFIX")]
    cands += [
        r"C:\Program Files\Tesseract-OCR\tessdata",
        r"C:\Program Files (x86)\Tesseract-OCR\tessdata",
        "/usr/share/tesseract-ocr/5/tessdata",
        "/usr/share/tesseract-ocr/4.00/tessdata",
        "/usr/share/tessdata",
        "/opt/homebrew/share/tessdata",
        "/usr/local/share/tessdata",
    ]
    for c in cands:
        if c and Path(c).is_dir():
            return str(c)
    return None


def collect_lines(textdict):
    lines = []
    for b in textdict.get("blocks", []):
        if b.get("type") != 0:
            continue
        for ln in b.get("lines", []):
            spans = [s for s in ln.get("spans", []) if s.get("text", "").strip()]
            if not spans:
                continue
            text = "".join(s["text"] for s in ln["spans"]).strip()
            main = max(spans, key=lambda s: s.get("size", 0))
            lines.append({
                "text": text,
                "bbox": ln["bbox"],
                "size": main.get("size", 12),
                "font": main.get("font", ""),
                "flags": int(main.get("flags", 0)),
                "color": int(main.get("color", 0)),
            })
    return lines


def extract_lines(page, use_ocr, lang, ocr_dpi):
    """Return (lines, source, from_ocr)."""
    lines = collect_lines(page.get_text("dict"))
    if lines:
        return lines, "embedded text", False
    if not use_ocr:
        return [], "no extractable text (OCR off)", False
    try:
        tp = page.get_textpage_ocr(language=lang or "eng", dpi=ocr_dpi, full=True,
                                   tessdata=find_tessdata())
        lines = collect_lines(page.get_text("dict", textpage=tp))
        return lines, ("OCR" if lines else "OCR found no text"), True
    except Exception as exc:
        return [], f"OCR failed: {exc}", False


def px_box(bbox, k, origin, shape, pad=0):
    h, w = shape[:2]
    x0, y0, x1, y1 = bbox
    return (max(0, int((x0 - origin[0]) * k) - pad),
            max(0, int((y0 - origin[1]) * k) - pad),
            min(w, int((x1 - origin[0]) * k) + pad + 1),
            min(h, int((y1 - origin[1]) * k) + pad + 1))


def ring_pixels(arr, box, pad=4):
    x0, y0, x1, y1 = box
    h, w = arr.shape[:2]
    e = arr[max(0, y0 - pad):min(h, y1 + pad), max(0, x0 - pad):min(w, x1 + pad)]
    if e.size == 0:
        return np.zeros((1, 3), dtype=int)
    t = 3
    ring = np.concatenate([e[:t].reshape(-1, 3), e[-t:].reshape(-1, 3),
                           e[:, :t].reshape(-1, 3), e[:, -t:].reshape(-1, 3)])
    return ring.astype(int)


def estimate_fg(arr, box, fallback):
    """Text colour = pixels furthest from the surrounding background colour."""
    x0, y0, x1, y1 = box
    crop = arr[y0:y1, x0:x1].reshape(-1, 3).astype(int)
    if crop.size == 0:
        return fallback
    bg = np.median(ring_pixels(arr, box), axis=0)
    dist = np.abs(crop - bg).sum(axis=1)
    if dist.max() < 60:  # no visible contrast; keep the PDF's own colour
        return fallback
    n = max(1, len(dist) // 10)
    idx = np.argsort(dist)[-n:]
    return tuple(int(v) for v in np.median(crop[idx], axis=0))


def clean_background(arr, boxes):
    """Remove text from the page image."""
    if not boxes:
        return arr
    out = arr.copy()
    mask = np.zeros(arr.shape[:2], np.uint8)
    for b in boxes:
        x0, y0, x1, y1 = b
        ring = ring_pixels(arr, b)
        if ring.std(axis=0).max() < 12 or cv2 is None:
            out[y0:y1, x0:x1] = np.median(ring, axis=0).astype(np.uint8)  # flat background
        else:
            mask[y0:y1, x0:x1] = 255  # gradient/photo: inpaint
    if mask.any() and cv2 is not None:
        out = cv2.inpaint(out, mask, 3, cv2.INPAINT_TELEA)
    return out


def add_line(slide, ln, rect, scale, rgb, from_ocr):
    x0, y0, x1, y1 = ln["bbox"]
    left = (x0 - rect.x0) * scale
    top = (y0 - rect.y0) * scale
    width = (x1 - x0) * scale * 1.10 + 0.05
    height = (y1 - y0) * scale + 0.02
    size_pdf = (y1 - y0) * OCR_FACTOR if from_ocr else ln["size"]
    size_pt = max(6, min(120, size_pdf * scale * 72))

    box = slide.shapes.add_textbox(Inches(max(0, left)), Inches(max(0, top)),
                                   Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = ln["text"]
    f = run.font
    font_name = ln["font"].split("+")[-1].split(",")[0].split("-")[0] if ln["font"] and not from_ocr else ""
    f.name = font_name or "Arial"
    f.size = Pt(size_pt)
    f.bold = bool(ln["flags"] & 16) or "bold" in ln["font"].lower()
    f.italic = bool(ln["flags"] & 2)
    f.color.rgb = RGBColor(*rgb)


# ---------------------------------------------------------------- converter
def convert_pdf(pdf_bytes, editable, use_ocr, lang, dpi, progress=None):
    report, problems = [], []
    k = dpi / 72.0
    with fitz.open(stream=pdf_bytes, filetype="pdf") as pdf:
        n = len(pdf)
        prs = Presentation()
        first = pdf[0].rect
        slide_w = 13.333
        slide_h = slide_w * first.height / first.width
        prs.slide_width, prs.slide_height = Inches(slide_w), Inches(slide_h)
        blank = prs.slide_layouts[6]

        for i, page in enumerate(pdf):
            rect = page.rect
            scale = slide_w / rect.width
            pix = page.get_pixmap(matrix=fitz.Matrix(k, k), alpha=False)
            arr = np.array(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
            slide = prs.slides.add_slide(blank)

            lines, source, from_ocr = [], "image only", False
            if editable:
                lines, source, from_ocr = extract_lines(page, use_ocr, lang, max(dpi, 250))

            colors = []
            if lines:
                boxes = [px_box(l["bbox"], k, (rect.x0, rect.y0), arr.shape, pad=3) for l in lines]
                for l, b in zip(lines, boxes):
                    c = l["color"]
                    colors.append(estimate_fg(arr, b, ((c >> 16) & 255, (c >> 8) & 255, c & 255)))
                arr = clean_background(arr, boxes)
            elif editable:
                problems.append(i + 1)

            buf = io.BytesIO()
            Image.fromarray(arr).save(buf, "PNG")
            buf.seek(0)
            slide.shapes.add_picture(buf, 0, 0, Inches(slide_w), Inches(slide_h))

            for l, rgb in zip(lines, colors):
                add_line(slide, l, rect, scale, rgb, from_ocr)

            report.append(f"Page {i + 1}: text boxes={len(lines)}; source={source}")
            if progress:
                progress((i + 1) / n, f"Converted {i + 1}/{n} pages")

        out = io.BytesIO()
        prs.save(out)
    return out.getvalue(), report, problems


# ---------------------------------------------------------------------- UI
st.set_page_config(page_title=APP_TITLE, page_icon="📄", layout="centered")
st.title(APP_TITLE)
st.write("Upload a NotebookLM PDF. Text becomes editable PowerPoint text; the rest of the "
         "slide stays as a clean background picture.")

tess_ok = shutil.which("tesseract") is not None
if tess_ok:
    st.success("Tesseract OCR found.")
else:
    st.error("Tesseract OCR was not found. NotebookLM PDFs contain no text layer, so without "
             "OCR the slides will be pictures only. Install Tesseract (Windows: UB Mannheim "
             "installer; Streamlit Cloud: put `tesseract-ocr` in packages.txt).")

uploaded = st.file_uploader("Upload your NotebookLM PDF", type=["pdf"])
mode = st.radio("Conversion mode",
                ["Editable slides (recommended)", "Image-only slides (exact look)"], index=0)
use_ocr = st.checkbox("Use OCR for pages with no text", value=True)
lang = st.text_input("OCR language code", value="eng", help="e.g. eng, eng+msa")
dpi = st.slider("Background image resolution (DPI)", 100, 300, 200, 25)
include_report = st.checkbox("Include conversion report (.txt)", value=True)

if uploaded:
    pdf_bytes = uploaded.getvalue()
    if not pdf_bytes.startswith(b"%PDF"):
        st.error("This file does not appear to be a valid PDF.")
        st.stop()
    if st.button("Convert PDF to PowerPoint", type="primary", use_container_width=True):
        try:
            bar = st.progress(0, text="Starting…")
            data, report, problems = convert_pdf(
                pdf_bytes, mode.startswith("Editable"), use_ocr, lang, dpi,
                progress=lambda v, t: bar.progress(v, text=t))
            st.session_state["pptx"] = data
            st.session_state["name"] = Path(uploaded.name).stem + "_editable.pptx"
            st.session_state["report"] = "\n".join(report)
            st.success(f"Conversion complete: {len(report)} slides.")
            if problems:
                st.warning(f"No text was recognised on pages {problems}. "
                           "Check Tesseract installation and the language code.")
        except Exception as exc:
            st.error(f"Conversion failed: {exc}")
            st.exception(exc)

if "pptx" in st.session_state:
    st.download_button(
        "Download PowerPoint (.pptx)", data=st.session_state["pptx"],
        file_name=st.session_state["name"],
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        type="primary", use_container_width=True)
    if include_report:
        st.download_button("Download conversion report (.txt)",
                           data=st.session_state["report"].encode("utf-8"),
                           file_name="conversion_report.txt", mime="text/plain",
                           use_container_width=True)
